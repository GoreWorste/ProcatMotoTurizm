# -*- coding: utf-8 -*-
"""Отчёт ПР3: титул из title_page_template.docx, оформление по методичке."""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
TITLE_TEMPLATE = ROOT / "tools" / "title_page_template.docx"
TITLE_FALLBACK = Path(r"c:\Users\hellt\Downloads\Копия Титульный лист.docx")
OUT = ROOT / "Otchet_Praktika_3_Formy_i_validaciya.docx"
UI = ROOT / "screenshots_report" / "pr3" / "ui"

FONT_NAME = "Times New Roman"
BLACK = RGBColor(0, 0, 0)
BODY_PT = 14
CAPTION_PT = 12
BODY_START_PARAGRAPH = 26  # с этого абзаца в шаблоне — чужой текст отчёта

_list_counter = 0


def _fmt_run(run, size_pt: int) -> None:
    run.font.name = FONT_NAME
    run.font.size = Pt(size_pt)
    run.font.color.rgb = BLACK
    run.bold = False
    run.italic = False
    run.underline = False


def _fmt_body_paragraph(paragraph, *, first_indent: bool = True) -> None:
    pf = paragraph.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = 1.5
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.first_line_indent = Cm(1.25) if first_indent else Pt(0)


def _set_paragraph_text(paragraph, text: str, size_pt: int, *, first_indent: bool = True) -> None:
    paragraph.clear()
    _fmt_body_paragraph(paragraph, first_indent=first_indent)
    run = paragraph.add_run(text)
    _fmt_run(run, size_pt)


def setup_normal_style(doc: Document) -> None:
    style = doc.styles["normal"] if "normal" in doc.styles else doc.styles["Normal"]
    style.font.name = FONT_NAME
    style.font.size = Pt(BODY_PT)
    style.font.color.rgb = BLACK
    style.font.bold = False
    style.font.italic = False
    pf = style.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = 1.5
    pf.first_line_indent = Cm(1.25)
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)


def trim_trailing_empty_paragraphs(doc: Document) -> None:
    while doc.paragraphs:
        last = doc.paragraphs[-1]
        if last.text.strip():
            break
        last._element.getparent().remove(last._element)


def strip_template_body(doc: Document) -> None:
    if len(doc.paragraphs) <= BODY_START_PARAGRAPH:
        return
    start = doc.paragraphs[BODY_START_PARAGRAPH]._element
    body = doc.element.body
    to_remove = []
    found = False
    for child in body:
        if child is start:
            found = True
        if found and child.tag != qn("w:sectPr"):
            to_remove.append(child)
    for el in to_remove:
        body.remove(el)


def _blacken_runs(paragraph) -> None:
    for r in paragraph.runs:
        r.font.color.rgb = BLACK
        r.bold = False
        r.italic = False
        r.underline = False


def format_title_page(doc: Document) -> None:
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    _blacken_runs(p)
    for p in doc.paragraphs:
        _blacken_runs(p)


def reset_list_num() -> None:
    global _list_counter
    _list_counter = 0


def add_h1(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    _set_paragraph_text(p, text, BODY_PT)


def add_h2(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    _set_paragraph_text(p, text, BODY_PT)


def add_p(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    _set_paragraph_text(p, text, BODY_PT)


def add_num(doc: Document, text: str) -> None:
    global _list_counter
    _list_counter += 1
    add_p(doc, f"{_list_counter}. {text}")


def add_bullet(doc: Document, text: str) -> None:
    add_p(doc, f"– {text}")


def add_code(doc: Document, line: str) -> None:
    p = doc.add_paragraph()
    _set_paragraph_text(p, line, BODY_PT)


def _format_table(table) -> None:
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                _fmt_body_paragraph(p, first_indent=False)
                if not p.runs:
                    t = p.text
                    p.clear()
                    if t:
                        _fmt_run(p.add_run(t), BODY_PT)
                for r in p.runs:
                    _fmt_run(r, BODY_PT)


def add_table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    try:
        table.style = "Table Grid"
    except KeyError:
        try:
            table.style = "Table Normal"
        except KeyError:
            pass
    for i, h in enumerate(headers):
        table.rows[0].cells[i].text = h
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            table.rows[ri + 1].cells[ci].text = val
    _format_table(table)


def add_image(
    doc: Document,
    path: Path,
    width_cm: float = 16.0,
    caption: str | None = None,
) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _fmt_body_paragraph(p, first_indent=False)
    if path.is_file():
        p.add_run().add_picture(str(path), width=Cm(width_cm))
    else:
        _fmt_run(p.add_run(f"[Снимок не найден: {path.name}]"), BODY_PT)
    if caption:
        cap = doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _fmt_body_paragraph(cap, first_indent=False)
        _fmt_run(cap.add_run(caption), CAPTION_PT)


def add_toc(doc: Document) -> None:
    lines = [
        "1. Введение",
        "2. Цель и задачи работы",
        "3. Используемые технологии",
        "4. Описание предметной области и сущностей",
        "5. Архитектура приложения",
        "6. Модели данных и объекты запросов",
        "7. Репозитории и персистентное хранение данных",
        "8. Управление состоянием (Provider)",
        "9. Пользовательский интерфейс",
        "10. Маршрутизация и параметры в URL",
        "11. Разбор ограничения при удалении связанного бренда",
        "12. Тестирование и проверка требований",
        "13. Примеры адресов с параметрами отбора",
        "14. Заключение",
        "Список использованных файлов проекта",
    ]
    add_h1(doc, "Содержание")
    for line in lines:
        add_p(doc, line)


def build_content(doc: Document) -> None:
    add_toc(doc)

    add_h1(doc, "1. Введение")
    add_p(
        doc,
        "В рамках итогового учебного проекта разрабатывается веб-приложение на Flutter "
        "для сервиса проката туристического и мототехнического оборудования. "
        "Практическая работа № 2 добавила каталоги со списками, поиском, фильтрацией "
        "и пагинацией; настоящая работа № 3 расширяет проект формами создания и "
        "редактирования, клиентской валидацией, отображением связей между сущностями "
        "и сохранением данных между перезагрузками страницы.",
    )
    add_p(
        doc,
        "Серверный backend не подключается: коллекции хранятся в браузере (localStorage "
        "через shared_preferences) в виде JSON. Репозитории выполняют CRUD, проверяют "
        "уникальность и ссылочную целостность; при повреждении JSON выполняется повторная "
        "загрузка демонстрационных данных (seed) и уведомление пользователя.",
    )

    add_h1(doc, "2. Цель и задачи работы")
    add_p(
        doc,
        "Цель работы — реализовать полный цикл работы с пятью сущностями предметной "
        "области через формы с валидацией, отобразить в интерфейсе три типа связей "
        "(многие-к-одному, многие-ко-многим, один-к-одному) и обеспечить персистентность "
        "данных без потери при обновлении страницы.",
    )
    add_p(doc, "Задачи:")
    reset_list_num()
    add_num(doc, "Расширить доменную модель: Equipment, Client, Category, Brand, Tag (+ RentalCard).")
    add_num(doc, "Заменить in-memory хранение на persistent-репозитории с shared_preferences.")
    add_num(doc, "Сохранить и распространить на все сущности возможности ПР2: поиск, сортировка, пагинация.")
    add_num(doc, "Реализовать экраны форм на Form / FormField с ошибками под полями.")
    add_num(doc, "Отобразить связи M2O (категория, бренд), M2M (теги), O2O (прокатный билет клиента).")
    add_num(doc, "Ввести уникальность инвентарного номера и email; ограничить удаление справочников со связями.")
    add_num(doc, "Обработать несохранённые изменения при уходе с формы (FormScaffold / PopScope).")
    add_num(doc, "Подготовить отчёт со скриншотами ключевых сценариев (валидация, связи, персистентность).")

    add_h1(doc, "3. Используемые технологии")
    add_bullet(doc, "Flutter (целевая платформа — Web);")
    add_bullet(doc, "Dart 3.x;")
    add_bullet(doc, "Управление состоянием: provider (ChangeNotifierProvider, context.watch / context.read);")
    add_bullet(doc, "Маршрутизация: go_router с usePathUrlStrategy() для «чистых» URL без #;")
    add_bullet(doc, "Персистентность: shared_preferences (localStorage в браузере);")
    add_bullet(doc, "Статический анализ: flutter analyze (в проекте — без замечаний);")
    add_bullet(doc, "Сборка: flutter build web --no-web-resources-cdn; запуск: scripts/run_web.ps1.")

    add_h1(doc, "4. Описание предметной области и сущностей")
    add_h2(doc, "4.1. Equipment (оборудование)")
    add_p(
        doc,
        "Единица инвентаря, сдаваемая в прокат. Поля: id, name, inventoryNumber (уникальный, "
        "формат EQ-…), categoryId, brandId, tagIds, purchaseYear, dailyRate, condition "
        "(«новое» | «хорошее» | «удовлетворительное»), unitsTotal, unitsAvailable, deletedAt.",
    )
    add_p(doc, "Справочники в данных:")
    add_bullet(doc, "Категории: Палатки, Велосипеды, Мотоциклы, Квадроциклы, Водный транспорт, Снегоходы;")
    add_bullet(doc, "Бренды: Tramp, Stels, Yamaha, Polaris, Trek, Sea-Doo, Buran;")
    add_bullet(doc, "Теги: зима, лето, для новичков, экстрим, водный.")
    add_p(doc, "В seed-данных — 22 единицы оборудования с реалистичными названиями.")

    add_h2(doc, "4.2. Client (клиент)")
    add_p(
        doc,
        "Поля: id, fullName, email (уникальный), phone, city, registeredAt, rentalCard "
        "(number, issuedAt), deletedAt. В начальных данных — 9 клиентов из городов Урала и Сибири.",
    )

    add_h2(doc, "4.3. Category (категория)")
    add_p(doc, "Справочник типов оборудования. Поля: id, name, deletedAt. Связь M2O с Equipment через categoryId.")

    add_h2(doc, "4.4. Brand (бренд)")
    add_p(
        doc,
        "Справочник производителей. Поля: id, name, deletedAt. Связь M2O с Equipment; "
        "физическое удаление запрещено при наличии активных единиц оборудования.",
    )

    add_h2(doc, "4.5. Tag (тег)")
    add_p(doc, "Справочник меток (сезон, уровень сложности). Поля: id, name, deletedAt. Связь M2M с Equipment через tagIds.")

    add_h2(doc, "4.6. Соответствие учебному заданию (библиотека → прокат)")
    add_table(
        doc,
        ["Учебная модель", "Проект", "Связь"],
        [
            ("Book", "Equipment", "M2O: Category, Brand; M2M: Tag"),
            ("Author", "Tag", "M2M с Equipment"),
            ("Genre", "Category", "M2O через categoryId"),
            ("Publisher", "Brand", "M2O; запрет hardDelete при связях"),
            ("Reader + билет", "Client + RentalCard", "O2O в форме клиента"),
        ],
    )
    add_p(doc, "Логическая схема: Client (1)—(1) RentalCard; Equipment (N)—(1) Category; Equipment (N)—(1) Brand; Equipment (N)—(M) Tag.")

    add_h1(doc, "5. Архитектура приложения")
    add_p(
        doc,
        "Приложение разделено на слои. Экраны списков не обращаются к репозиториям напрямую — "
        "через ChangeNotifier (EquipmentListNotifier, ClientListNotifier и аналоги для справочников). "
        "Экраны форм вызывают методы репозиториев после успешной валидации FormState.",
    )
    add_p(
        doc,
        "Схема потока данных: форма → validators + FormField → create/update в репозитории → "
        "AppDataStore.persist* → shared_preferences → при возврате на список notifier.load().",
    )
    add_table(
        doc,
        ["Слой", "Каталог", "Назначение"],
        [
            ("Модели", "lib/models/", "Сущности, RentalCard, PageResult, Query, LoadStatus"),
            ("Репозитории", "lib/repositories/", "Интерфейсы, AppDataStore, Persistent* реализации"),
            ("Состояние", "lib/state/", "ChangeNotifier для списков и выборки"),
            ("Экраны", "lib/screens/", "Списки, детали, формы equipment/client/справочников"),
            ("Виджеты", "lib/widgets/", "EntityTable, FormScaffold, MultiIdChipsField, поиск, пагинация"),
            ("Маршруты", "lib/routing/", "GoRouter, парсинг query-параметров"),
            ("Валидация", "lib/core/", "Переиспользуемые validators.dart"),
        ],
    )
    add_image(doc, UI / "00_schema_relations.png", caption="Рисунок 1 - Схема связей сущностей")

    add_h1(doc, "6. Модели данных и объекты запросов")
    add_p(
        doc,
        "Классы Equipment, Client, Category, Brand, Tag — неизменяемые (immutable), с copyWith, "
        "toJson и fromJson. Для deletedAt используется флаг clearDeletedAt в copyWith. "
        "Equipment хранит связи как categoryId, brandId и список tagIds.",
    )
    add_p(
        doc,
        "RentalCard — вложенная модель клиента (номер билета и дата выдачи), сериализуется "
        "внутри объекта Client в JSON clients_v1.",
    )
    add_p(
        doc,
        "PageResult<T> содержит items, page, size, total; геттеры totalPages, hasPrevious, hasNext; "
        "конструктор PageResult.empty() для начального состояния.",
    )
    add_p(
        doc,
        "EquipmentQuery, ClientQuery и NamedEntityQuery — объекты отбора для списков. "
        "При изменении условий поиска номер страницы сбрасывается на 1 (как в ПР2).",
    )
    add_p(doc, "Фильтры оборудования (комбинируются одновременно):")
    add_bullet(doc, "categoryId — категория;")
    add_bullet(doc, "brandId — бренд;")
    add_bullet(doc, "dailyRateFrom / dailyRateTo — диапазон суточного тарифа;")
    add_bullet(doc, "yearFrom / yearTo — диапазон года покупки;")
    add_bullet(doc, "includeDeleted — показ логически удалённых записей.")
    add_p(doc, "Правила валидации форм (кратко):")
    add_table(
        doc,
        ["Поле", "Правила"],
        [
            ("inventoryNumber", "формат EQ-…, уникальность"),
            ("email клиента", "формат, уникальность"),
            ("теги оборудования", "минимум один"),
            ("unitsAvailable", "не больше unitsTotal"),
            ("имя справочника", "2–80 символов, обязательность"),
        ],
    )

    add_h1(doc, "7. Репозитории и персистентное хранение данных")
    add_p(
        doc,
        "AppDataStore загружает при старте ключи equipment_v1, clients_v1, categories_v1, "
        "brands_v1, tags_v1, meta_v1. Каждый репозиторий держит кэш списка в памяти и после "
        "изменений вызывает persist*. Методы интерфейсов: find, findById, create, update, "
        "softDelete, hardDelete, restore, deleteMany (где применимо).",
    )
    add_p(doc, "Метод find() в persistent-реализации:")
    reset_list_num()
    add_num(doc, "Искусственная задержка ~250 мс (имитация сети, индикатор загрузки).")
    add_num(doc, "Фильтр includeDeleted / скрытие записей с deletedAt.")
    add_num(doc, "Поиск по текстовым полям сущности (имя, инвентарный номер, email и т.д.).")
    add_num(doc, "Сортировка по настраиваемому полю и направлению (не менее трёх полей на сущность).")
    add_num(doc, "Пагинация через sublist; корректный расчёт границ страницы.")
    add_num(doc, "При create/update оборудования и клиентов — проверка уникальности ключевых полей в репозитории.")
    add_p(
        doc,
        "При повреждении JSON в localStorage AppDataStore выполняет reseed и передаёт сообщение "
        "в UI (SnackBar), чтобы пользователь знал о сбросе данных.",
    )

    add_h1(doc, "8. Управление состоянием (Provider)")
    add_p(
        doc,
        "Для каждой сущности списка — свой ChangeNotifier: query, result (PageResult), "
        "status (LoadStatus: idle, loading, success, error), error, selected (Set<int>).",
    )
    add_p(
        doc,
        "Ключевые методы: load(), applyQuery(next), toggleSelection(id), операции удаления "
        "и восстановления — как в ПР2. После успешного сохранения формы выполняется "
        "context.pop() и при необходимости load() на целевом списке.",
    )
    add_p(
        doc,
        "На формах локальное состояние (выбранные id тегов, выпадающие списки) обновляется "
        "через setState; данные каталога для выпадающих списков читаются через context.read.",
    )

    add_h1(doc, "9. Пользовательский интерфейс")
    add_h2(doc, "9.1. Обобщённая таблица EntityTable<T>")
    add_p(
        doc,
        "Виджет lib/widgets/entity_table.dart принимает список TableColumnSpec<T> (label, sortField, "
        "numeric, build), данные, idOf, выделение, параметры сортировки и колбэки — без изменений "
        "по сравнению с ПР2, используется для всех пяти каталогов.",
    )

    add_h2(doc, "9.2. Адаптивная вёрстка")
    add_p(
        doc,
        "При ширине окна менее 600 px вместо таблицы отображается ListView с Card для каждой "
        "записи (те же действия: выделение, открытие карточки, переход к редактированию, удаление).",
    )

    add_h2(doc, "9.3. Состояния экрана")
    add_table(
        doc,
        ["Состояние", "Отображение"],
        [
            ("loading", "CircularProgressIndicator по центру"),
            ("success (есть данные)", "Таблица или список карточек"),
            ("success (пусто)", "ListEmptyView — иконка и текст «Ничего не найдено…»"),
            ("error", "ListErrorView — иконка ошибки, текст, кнопка «Повторить»"),
        ],
    )

    add_h2(doc, "9.4. Поиск, фильтры, пагинация")
    add_bullet(doc, "DebouncedSearchField — задержка 350 мс после ввода;")
    add_bullet(doc, "Сортировка по клику на заголовок колонки со стрелкой направления;")
    add_bullet(doc, "Пагинация: размер страницы 10 / 25 / 50, кнопки первая / пред / след / последняя, отображение page и total.")

    add_h2(doc, "9.5. Формы создания и редактирования")
    add_p(
        doc,
        "Equipment: поля инвентаря + блок «Связи с другими сущностями» (DropdownButtonFormField "
        "для категории и бренда, MultiIdChipsField для тегов). При смене категории список "
        "брендов фильтруется; недопустимый brandId сбрасывается.",
    )
    add_p(
        doc,
        "Client: контактные поля + секция прокатного билета (O2O). CatalogEntityFormScreen — "
        "единая форма для Category, Brand, Tag.",
    )
    add_p(
        doc,
        "FormScaffold: кнопки «Сохранить» / «Отмена», PopScope с диалогом при несохранённых изменениях.",
    )

    add_h2(doc, "9.6. Маршруты экранов")
    add_bullet(doc, "/equipment, /equipment/new, /equipment/:id, /equipment/:id/edit;")
    add_bullet(doc, "/clients, /clients/new, /clients/:id, /clients/:id/edit;")
    add_bullet(doc, "/categories, /brands, /tags — списки и формы /new, /:id/edit.")

    add_h2(doc, "9.7. Скриншоты интерфейса")
    ui_items = [
        ("01_equipment_form_relations.png", "Рисунок 2 - Форма оборудования: связи M2O/M2M"),
        ("02_equipment_validation_errors.png", "Рисунок 3 - Ошибки валидации на форме"),
        ("03_equipment_duplicate_inventory.png", "Рисунок 4 - Дубликат инвентарного номера"),
        ("04_client_form_rental_card.png", "Рисунок 5 - Клиент и прокатный билет (O2O)"),
        ("05_brand_delete_blocked.png", "Рисунок 6 - Удаление бренда со связанным оборудованием"),
        ("06_persistence_before_reload.png", "Рисунок 7 - Данные до перезагрузки страницы"),
        ("07_persistence_after_reload.png", "Рисунок 8 - Данные после F5"),
        ("08_categories_search_pagination.png", "Рисунок 9 - Справочник: поиск и пагинация"),
        ("09_equipment_list_pr2_features.png", "Рисунок 10 - Список оборудования с фильтрами"),
    ]
    for fname, cap in ui_items:
        add_image(doc, UI / fname, caption=cap)

    add_h1(doc, "10. Маршрутизация и параметры в URL")
    add_p(
        doc,
        "В main.dart настроен GoRouter. Условия отбора списков кодируются в query-параметрах, "
        "например: /equipment?search=палатка&categoryId=1&sort=dailyRate,desc&page=1&size=10. "
        "Формы открываются по путям /new и /:id/edit без потери возможности вернуться к списку.",
    )
    add_p(
        doc,
        "При открытии URL в новой вкладке экран списка в didChangeDependencies парсит параметры "
        "(lib/routing/query_params.dart) и вызывает applyQuery — поведение ПР2 сохранено.",
    )

    add_h1(doc, "11. Разбор ограничения при удалении связанного бренда")
    add_p(
        doc,
        "По требованиям практики издательство (в проекте — Brand) нельзя физически удалить, "
        "если на него ссылается хотя бы одна активная единица оборудования. Проверка выполняется "
        "в репозитории до изменения списка; UI перехватывает ReferenceInUseException и показывает SnackBar.",
    )
    add_p(doc, "Фрагмент lib/repositories/persistent_brand_repository.dart:")
    add_code(doc, "Future<void> hardDelete(int id) async {")
    add_code(doc, "  final count = countEquipmentLinks(id);")
    add_code(doc, "  if (count > 0) {")
    add_code(doc, "    throw ReferenceInUseException(")
    add_code(doc, "      'Нельзя удалить бренд: связано единиц оборудования — $count',")
    add_code(doc, "      count: count,")
    add_code(doc, "    );")
    add_code(doc, "  }")
    add_code(doc, "  _items.removeWhere((b) => b.id == id);")
    add_code(doc, "  await _store.persistBrands();")
    add_code(doc, "}")
    add_p(
        doc,
        "Аналогичная логика применяется для Category. Таким образом, ссылочная целостность "
        "поддерживается на уровне домена, а не только подсказкой в интерфейсе.",
    )

    add_h1(doc, "12. Тестирование и проверка требований")
    add_p(doc, "Выполненные проверки:")
    add_bullet(doc, "flutter analyze — замечаний нет;")
    add_bullet(doc, "flutter build web --no-web-resources-cdn — сборка успешна;")
    add_bullet(doc, "Ручная проверка: сохранение формы и F5 — данные в localStorage сохраняются;")
    add_bullet(doc, "Ручная проверка: дубликат inventoryNumber и email — ошибка под полем;")
    add_bullet(doc, "Ручная проверка: пустая форма — несколько ошибок одновременно;")
    add_bullet(doc, "Ручная проверка: удаление бренда со связями — SnackBar с количеством;")
    add_bullet(doc, "Ручная проверка: функции ПР2 на equipment, clients, categories, brands, tags.")
    add_table(
        doc,
        ["Типичная ошибка (методичка)", "Как избежано в проекте"],
        [
            ("Связи не видны в UI", "Отдельный блок на форме оборудования + подписи M2O/M2M"),
            ("Уникальность только в форме", "Дополнительная проверка в репозитории при save"),
            ("Данные пропадают после F5", "persist после каждого create/update/delete"),
            ("Можно уйти с несохранённой формой", "FormScaffold и PopScope с диалогом"),
            ("M2M не валидируется", "FormField на MultiIdChipsField, минимум один тег"),
            ("Provider не найден", "MultiProvider в main.dart"),
            ("Поиск без debounce", "DebouncedSearchField 350 мс"),
        ],
    )
    add_p(doc, "Команды запуска:")
    add_p(doc, "cd procatMotoTurizm")
    add_p(doc, "flutter pub get")
    add_p(doc, ".\\scripts\\run_web.ps1")

    add_h1(doc, "13. Примеры адресов с параметрами отбора")
    add_p(doc, "Базовый хост при локальном запуске: http://localhost:<порт>.")
    add_p(doc, "Оборудование:")
    add_bullet(doc, "/equipment/new")
    add_bullet(doc, "/equipment?search=палатка&categoryId=1&sort=dailyRate,desc&page=1&size=10")
    add_bullet(doc, "/equipment?categoryId=4&brandId=2&dailyRateFrom=3000&dailyRateTo=5000&sort=name,asc&page=1&size=25")
    add_bullet(doc, "/equipment?includeDeleted=true&sort=name,asc&page=1&size=10")
    add_p(doc, "Клиенты:")
    add_bullet(doc, "/clients/new")
    add_bullet(doc, "/clients?search=Иванов&city=Екатеринбург&sort=fullName,asc&page=1&size=10")
    add_bullet(doc, "/clients?sort=registeredAt,desc&page=1&size=25")
    add_p(doc, "Справочники:")
    add_bullet(doc, "/categories?search=мото&page=1&size=10")
    add_bullet(doc, "/brands")
    add_bullet(doc, "/tags")

    add_h1(doc, "14. Заключение")
    add_p(
        doc,
        "В ходе практической работы № 3 к Flutter Web-приложению проката добавлены пять "
        "персистентных сущностей с полным CRUD, формы с валидацией, отображение связей "
        "M2O, M2M и O2O, уникальность ключевых полей и ограничения при удалении справочников. "
        "Сохранены и расширены на все каталоги возможности ПР2 (поиск, фильтры, сортировка, "
        "пагинация, синхронизация с URL). Данные сохраняются между сессиями в localStorage.",
    )

    add_h1(doc, "Список основных файлов проекта")
    add_bullet(doc, "lib/main.dart")
    add_bullet(doc, "lib/core/validators.dart")
    add_bullet(doc, "lib/models/* (equipment, client, category, brand, tag, rental_card, queries, page_result)")
    add_bullet(doc, "lib/repositories/app_data_store.dart, persistent_*_repository.dart")
    add_bullet(doc, "lib/state/*_list_notifier.dart")
    add_bullet(doc, "lib/routing/app_router.dart, query_params.dart")
    add_bullet(
        doc,
        "lib/screens/* (списки, детали, equipment_form_screen, client_form_screen, catalog_entity_form_screen)",
    )
    add_bullet(
        doc,
        "lib/widgets/* (entity_table, form_scaffold, multi_id_chips_field, pagination, search, состояния)",
    )
    add_bullet(doc, "scripts/run_web.ps1")


def main() -> None:
    title_path = TITLE_TEMPLATE if TITLE_TEMPLATE.is_file() else TITLE_FALLBACK
    if not title_path.is_file():
        raise SystemExit(f"Нет титульного листа: {TITLE_TEMPLATE}")

    doc = Document(str(title_path))
    strip_template_body(doc)
    trim_trailing_empty_paragraphs(doc)
    setup_normal_style(doc)
    format_title_page(doc)
    doc.add_page_break()
    add_h1(
        doc,
        "Практическая работа № 3. Формы, валидация и связанные сущности",
    )
    build_content(doc)

    try:
        doc.save(OUT)
        print(OUT)
    except PermissionError:
        alt = ROOT / "Otchet_Praktika_3_Formy_i_validaciya.generated.docx"
        doc.save(alt)
        print(alt)
        print("(Закройте открытый DOCX и перезапустите скрипт.)")


if __name__ == "__main__":
    main()

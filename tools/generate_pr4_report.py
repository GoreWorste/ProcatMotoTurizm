# -*- coding: utf-8 -*-
"""Отчёт ПР4: REST API, Dio, оформление как ПР3."""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
TITLE_TEMPLATE = ROOT / "tools" / "title_page_template.docx"
TITLE_FALLBACK = Path(r"c:\Users\hellt\Downloads\Копия Титульный лист.docx")
OUT = ROOT / "Otchet_Praktika_4_REST_API.docx"
UI = ROOT / "screenshots_report" / "pr4" / "ui"
CODE = ROOT / "screenshots_report" / "pr4" / "code"

FONT_NAME = "Times New Roman"
BLACK = RGBColor(0, 0, 0)
BODY_PT = 14
CAPTION_PT = 12
BODY_START_PARAGRAPH = 26  # с этого абзаца в шаблоне — чужой текст отчёта

_list_counter = 0


def _fmt_run(run, size_pt: int) -> None:
    run.font.name = FONT_NAME
    run.font.size = Pt(size_pt)
    run.font.color.rgb = BLACK
    run.bold = False
    run.italic = False
    run.underline = False


def _fmt_body_paragraph(paragraph, *, first_indent: bool = True) -> None:
    pf = paragraph.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = 1.5
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.first_line_indent = Cm(1.25) if first_indent else Pt(0)


def _set_paragraph_text(paragraph, text: str, size_pt: int, *, first_indent: bool = True) -> None:
    paragraph.clear()
    _fmt_body_paragraph(paragraph, first_indent=first_indent)
    run = paragraph.add_run(text)
    _fmt_run(run, size_pt)


def setup_normal_style(doc: Document) -> None:
    style = doc.styles["normal"] if "normal" in doc.styles else doc.styles["Normal"]
    style.font.name = FONT_NAME
    style.font.size = Pt(BODY_PT)
    style.font.color.rgb = BLACK
    style.font.bold = False
    style.font.italic = False
    pf = style.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = 1.5
    pf.first_line_indent = Cm(1.25)
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)


def trim_trailing_empty_paragraphs(doc: Document) -> None:
    while doc.paragraphs:
        last = doc.paragraphs[-1]
        if last.text.strip():
            break
        last._element.getparent().remove(last._element)


def strip_template_body(doc: Document) -> None:
    if len(doc.paragraphs) <= BODY_START_PARAGRAPH:
        return
    start = doc.paragraphs[BODY_START_PARAGRAPH]._element
    body = doc.element.body
    to_remove = []
    found = False
    for child in body:
        if child is start:
            found = True
        if found and child.tag != qn("w:sectPr"):
            to_remove.append(child)
    for el in to_remove:
        body.remove(el)


def _blacken_runs(paragraph) -> None:
    for r in paragraph.runs:
        r.font.color.rgb = BLACK
        r.bold = False
        r.italic = False
        r.underline = False


def format_title_page(doc: Document) -> None:
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    _blacken_runs(p)
    for p in doc.paragraphs:
        _blacken_runs(p)


def reset_list_num() -> None:
    global _list_counter
    _list_counter = 0


def add_h1(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    _set_paragraph_text(p, text, BODY_PT)


def add_h2(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    _set_paragraph_text(p, text, BODY_PT)


def add_p(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    _set_paragraph_text(p, text, BODY_PT)


def add_num(doc: Document, text: str) -> None:
    global _list_counter
    _list_counter += 1
    add_p(doc, f"{_list_counter}. {text}")


def add_bullet(doc: Document, text: str) -> None:
    add_p(doc, f"– {text}")


def add_code(doc: Document, line: str) -> None:
    p = doc.add_paragraph()
    _set_paragraph_text(p, line, BODY_PT)


def _format_table(table) -> None:
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                _fmt_body_paragraph(p, first_indent=False)
                if not p.runs:
                    t = p.text
                    p.clear()
                    if t:
                        _fmt_run(p.add_run(t), BODY_PT)
                for r in p.runs:
                    _fmt_run(r, BODY_PT)


def add_table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    try:
        table.style = "Table Grid"
    except KeyError:
        try:
            table.style = "Table Normal"
        except KeyError:
            pass
    for i, h in enumerate(headers):
        table.rows[0].cells[i].text = h
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            table.rows[ri + 1].cells[ci].text = val
    _format_table(table)


def add_image(
    doc: Document,
    path: Path,
    width_cm: float = 16.0,
    caption: str | None = None,
) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _fmt_body_paragraph(p, first_indent=False)
    if path.is_file():
        p.add_run().add_picture(str(path), width=Cm(width_cm))
    else:
        _fmt_run(p.add_run(f"[Снимок не найден: {path.name}]"), BODY_PT)
    if caption:
        cap = doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _fmt_body_paragraph(cap, first_indent=False)
        _fmt_run(cap.add_run(caption), CAPTION_PT)


def add_toc(doc: Document) -> None:
    lines = [
        "1. Введение",
        "2. Цель и задачи работы",
        "3. Используемые технологии",
        "4. Контракт учебного API",
        "5. Архитектура: замена реализации репозитория",
        "6. Настройка Dio и исключения предметной области",
        "7. API-репозитории и guard",
        "8. Состояния сетевого слоя в UI",
        "9. Ошибки 422 и 409 в формах",
        "10. Диагностика CORS и сетевых сбоев",
        "11. Тестирование",
        "12. Заключение",
        "Список использованных файлов проекта",
    ]
    add_h1(doc, "Содержание")
    for line in lines:
        add_p(doc, line)


def build_content(doc: Document) -> None:
    add_toc(doc)

    add_h1(doc, "1. Введение")
    add_p(
        doc,
        "Учебное веб-приложение «Прокат мототехники» на Flutter в предыдущих работах "
        "получило каталоги со списками, поиском и пагинацией (ПР2), а также формы, "
        "валидацию и локальное хранение в браузере (ПР3). Практическая работа № 4 "
        "переводит приложение на клиент–серверную схему: данные пяти сущностей "
        "(Equipment, Client, Category, Brand, Tag) загружаются и изменяются через "
        "учебный REST API, реализованный на Node.js.",
    )
    add_p(
        doc,
        "Интерфейс и доменные интерфейсы репозиториев сохранены: меняется только "
        "реализация слоя доступа к данным (Dio вместо shared_preferences). "
        "Для отладки без сервера оставлен режим USE_API=false (локальный ПР3).",
    )

    add_h1(doc, "2. Цель и задачи работы")
    add_p(
        doc,
        "Цель — подключить REST API к существующему Flutter Web-приложению, "
        "обеспечить обработку сетевых сбоев и ответов сервера с кодами 422 и 409, "
        "сохранив функции ПР2–ПР3 при работе через HTTP.",
    )
    add_p(doc, "Задачи:")
    reset_list_num()
    add_num(doc, "Описать и реализовать учебный контракт API (CRUD, пагинация, фильтры).")
    add_num(doc, "Настроить HTTP-клиент Dio (baseUrl, таймауты, interceptors, повтор GET).")
    add_num(doc, "Ввести типизированные исключения и функции guard / describeError.")
    add_num(doc, "Реализовать Api*Repository для всех сущностей поверх тех же интерфейсов.")
    add_num(doc, "Отобразить в UI состояния загрузки, ошибки и отмену устаревших запросов.")
    add_num(doc, "Сопоставить ошибки 422 полям форм, 409 — конфликтам удаления справочников.")
    add_num(doc, "Обеспечить CORS при раздельном запуске клиента (порт 5555) и API (8080).")
    add_num(doc, "Подготовить отчёт со скриншотами сценариев работы через API.")

    add_h1(doc, "3. Используемые технологии")
    add_bullet(doc, "Flutter Web, Dart 3, provider, go_router;")
    add_bullet(doc, "HTTP-клиент: пакет dio 5.x;")
    add_bullet(doc, "Конфигурация сборки: --dart-define=USE_API, --dart-define=API_BASE_URL;")
    add_bullet(doc, "Учебный backend: api/mock-server.js, данные api/db.json;")
    add_bullet(doc, "Экспорт seed из локального хранилища: dart run tools/export_api_seed.dart;")
    add_bullet(doc, "Тесты: flutter test test/api_equipment_repository_test.dart (http_mock_adapter).")

    add_h1(doc, "4. Контракт учебного API")
    add_p(
        doc,
        "Базовый URL: http://localhost:8080/api. Служебные возможности: GET /__health, "
        "искусственная задержка ?__delay=мс и учебная ошибка ?__fail=код для демонстрации "
        "сетевого слоя. Полное описание — в файле api/КОНТРАКТ-API.md.",
    )
    add_table(
        doc,
        ["Ресурс", "Основные операции"],
        [
            ("/equipment", "GET (фильтры, sort, page, size), POST, PUT, DELETE, restore, bulk-delete"),
            ("/clients", "Аналогично; 422 при дубликате email"),
            ("/categories, /brands, /tags", "CRUD, bulk-delete, restore; 409 при hard delete со связями"),
        ],
    )
    add_p(doc, "Формат страницы списка:")
    add_p(doc, '{ "items": [...], "page": 1, "size": 10, "total": N }')
    add_image(
        doc,
        UI / "00_api_health.png",
        caption="Рисунок 1 — Проверка доступности API (GET /api/__health)",
    )

    add_h1(doc, "5. Архитектура: замена реализации репозитория")
    add_p(
        doc,
        "В lib/main.dart по флагу useApiBackend (из USE_API) создаётся либо связка "
        "ApiEquipmentRepository, ApiClientRepository и Api* для справочников на общем Dio, "
        "либо прежние Persistent*Repository и AppDataStore. Экраны и ChangeNotifier "
        "зависят только от абстрактных интерфейсов в lib/repositories/*_repository.dart.",
    )
    add_table(
        doc,
        ["Компонент", "Назначение"],
        [
            ("lib/core/config.dart", "apiBaseUrl, useApiBackend"),
            ("lib/core/api_client.dart", "buildDio(), retry GET"),
            ("lib/repositories/api/*", "HTTP-реализации CRUD и разбор PageResult"),
            ("lib/state/*_notifier.dart", "load(), CancelToken, describeError в UI"),
            ("api/mock-server.js", "Учебный сервер с CORS и правилами 422/409"),
        ],
    )
    add_image(
        doc,
        CODE / "main_api_wiring.dart.png",
        caption="Рисунок 2 — Подключение API-репозиториев в main.dart",
    )

    add_h1(doc, "6. Настройка Dio и исключения предметной области")
    add_p(
        doc,
        "Dio создаётся с connectTimeout 10 с и receiveTimeout 15 с. Interceptor "
        "преобразует ответы с кодом ≥400 в DioException с уже разобранным ApiException "
        "(mapHttpError). Для GET добавлен _ReadRetryInterceptor — до трёх повторов "
        "при connectionTimeout / connectionError с паузой 300·n мс.",
    )
    add_image(doc, CODE / "api_client.dart.png", caption="Рисунок 3 — Конфигурация Dio и interceptor")
    add_p(
        doc,
        "Классы NetworkException, ValidationException (с полем errors), ConflictException, "
        "NotFoundException и др. позволяют экранам показывать осмысленные сообщения. "
        "Функция guard оборачивает вызовы репозитория; describeError унифицирует вывод в SnackBar и ListErrorView.",
    )
    add_image(
        doc,
        CODE / "api_exceptions.dart.png",
        caption="Рисунок 4 — Маппинг HTTP-кодов и mapDioError (в т.ч. подсказка про CORS)",
    )

    add_h1(doc, "7. API-репозитории и guard")
    add_p(
        doc,
        "ApiEquipmentRepository выполняет GET с queryParameters из EquipmentQuery, "
        "разбирает JSON через parsePage и fromJson модели. create/update передают тело "
        "запроса; при 422 пробрасывается ValidationException с ключами полей, совпадающими "
        "с именами в форме Flutter (inventoryNumber, email и т.д.).",
    )
    add_image(
        doc,
        CODE / "api_equipment_find.dart.png",
        caption="Рисунок 5 — Пример запроса списка оборудования через Dio",
    )
    add_p(
        doc,
        "Справочники (категории, бренды, теги) вынесены в api_catalog_repositories.dart. "
        "После изменений справочников CatalogNotifier.invalidateCache() сбрасывает кэш, "
        "чтобы формы и фильтры получали актуальные id.",
    )

    add_h1(doc, "8. Состояния сетевого слоя в UI")
    add_p(
        doc,
        "Списки оборудования и клиентов при load() переводятся в LoadStatus.loading, "
        "по завершении — success или error с текстом describeError. При новом запросе "
        "отменяется предыдущий CancelToken (актуально для быстрого ввода в поиске). "
        "FormScaffold поддерживает флаг submitting на время POST/PUT.",
    )
    add_image(
        doc,
        UI / "01_equipment_list.png",
        caption="Рисунок 6 — Список оборудования, данные с GET /api/equipment (пагинация «всего 22»)",
    )
    add_image(
        doc,
        UI / "02_equipment_filters.png",
        caption="Рисунок 7 — Фильтр categoryId=1 в URL и ответ API (всего 4 записи категории «Палатки»)",
    )
    add_image(
        doc,
        UI / "03_clients_list.png",
        caption="Рисунок 8 — Список клиентов с GET /api/clients",
    )
    add_image(
        doc,
        UI / "08_brands_list.png",
        caption="Рисунок 9 — Справочник брендов (GET /api/brands)",
    )
    add_p(
        doc,
        "При недоступности сервера список переходит в ListErrorView с кнопкой «Повторить» "
        "(сообщение из NetworkException). На рисунке 10 показан интерфейс после остановки "
        "mock-сервера и обновления страницы.",
    )
    add_image(
        doc,
        UI / "06_list_network_error.png",
        caption="Рисунок 10 — Сбой загрузки каталога при остановленном API",
    )

    add_h1(doc, "9. Ошибки 422 и 409 в формах и справочниках")
    add_p(
        doc,
        "Сервер при дубликате инвентарного номера возвращает 422 и объект errors. "
        "EquipmentFormScreen записывает их в _serverErrors и показывает под полем "
        "«Инвентарный номер». Аналогично для email клиента.",
    )
    add_image(
        doc,
        UI / "04_api_422_duplicate.png",
        caption="Рисунок 11 — Ответ POST /api/equipment с кодом 422 (EQ-1001 уже существует)",
    )
    add_image(
        doc,
        UI / "07_equipment_form_new.png",
        caption="Рисунок 12 — Форма создания оборудования (данные отправляются POST /api/equipment)",
    )
    add_p(
        doc,
        "Физическое удаление бренда или категории при связанном оборудовании "
        "возвращает 409. В UI ConflictException отображается через SnackBar "
        "(экран детали справочника и списки).",
    )
    add_image(
        doc,
        UI / "05_api_409_brand.png",
        caption="Рисунок 13 — Ответ DELETE /api/brands/1?hard=true с кодом 409",
    )

    add_h1(doc, "10. Диагностика CORS и сетевых сбоев")
    add_p(
        doc,
        "Клиент по умолчанию открывается на http://localhost:5555, API — на порту 8080. "
        "В mock-server.js заголовок Access-Control-Allow-Origin задаётся параметром --origin "
        "(scripts/run_api_server.ps1 передаёт http://localhost:5555). При несовпадении "
        "origin браузер блокирует ответ; mapDioError для connectionError содержит "
        "подсказку проверить консоль на CORS.",
    )
    add_p(doc, "Порядок запуска для демонстрации:")
    add_bullet(doc, ".\\scripts\\run_api_server.ps1")
    add_bullet(doc, ".\\scripts\\run_web.ps1")
    add_bullet(doc, "Проверка: http://localhost:8080/api/__health")

    add_h1(doc, "11. Тестирование")
    add_p(doc, "Автоматизированные проверки слоя API:")
    add_bullet(doc, "разбор постраничного ответа find;")
    add_bullet(doc, "create с 422 → ValidationException;")
    add_bullet(doc, "mapHttpError для 404;")
    add_bullet(doc, "mapDioError для connectionError и сохранение ValidationException в error.")
    add_p(doc, "Команда: flutter test test/api_equipment_repository_test.dart")
    add_p(
        doc,
        "Виджет-тест App builds выполняется только при USE_API=false, чтобы не зависеть "
        "от запущенного Node-сервера в CI.",
    )

    add_h1(doc, "12. Заключение")
    add_p(
        doc,
        "В практической работе № 4 приложение проката переведено на работу с учебным "
        "REST API: реализованы контракт и mock-сервер, слой Dio с повтором чтения и "
        "маппингом ошибок, HTTP-репозитории для пяти сущностей. Сохранены поиск, "
        "фильтры, пагинация и формы ПР2–ПР3; ошибки валидации и конфликты удаления "
        "обрабатываются по кодам 422 и 409. Переключатель USE_API позволяет "
        "сравнивать локальный и сетевой режимы.",
    )

    add_h1(doc, "Список основных файлов проекта")
    add_bullet(doc, "lib/core/config.dart, api_client.dart, api_exceptions.dart")
    add_bullet(doc, "lib/repositories/api/*.dart")
    add_bullet(doc, "lib/main.dart")
    add_bullet(doc, "api/mock-server.js, api/db.json, api/КОНТРАКТ-API.md")
    add_bullet(doc, "scripts/run_api_server.ps1, scripts/run_web.ps1")
    add_bullet(doc, "test/api_equipment_repository_test.dart")
    add_bullet(doc, "tools/capture_pr4_screenshots.py, tools/generate_pr4_report.py")

def main() -> None:
    title_path = TITLE_TEMPLATE if TITLE_TEMPLATE.is_file() else TITLE_FALLBACK
    if not title_path.is_file():
        raise SystemExit(f"Нет титульного листа: {TITLE_TEMPLATE}")

    doc = Document(str(title_path))
    strip_template_body(doc)
    trim_trailing_empty_paragraphs(doc)
    setup_normal_style(doc)
    format_title_page(doc)
    doc.add_page_break()
    add_h1(
        doc,
        "Практическая работа № 4. Подключение REST API",
    )
    build_content(doc)

    try:
        doc.save(OUT)
        print(OUT)
    except PermissionError:
        alt = ROOT / "Otchet_Praktika_4_REST_API.generated.docx"
        doc.save(alt)
        print(alt)
        print("(Закройте открытый DOCX и перезапустите скрипт.)")


if __name__ == "__main__":
    main()

# -*- coding: utf-8 -*-
"""Отчёт: титульник, обязательные разделы, скриншоты UI и кода (PNG)."""

from __future__ import annotations

from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

ROOT = Path(__file__).resolve().parents[1]
TITLE_TEMPLATE = Path(r"c:\Users\hellt\Downloads\Копия Титульный лист.docx")
SCREENSHOTS = ROOT / "screenshots_report"
CODE_DIR = SCREENSHOTS / "code"
OUTPUT = ROOT / "Отчет_Практика_2_Списки_и_пагинация.docx"
GITHUB_DEV = "https://github.com/GoreWorste/ProcatMotoTurizm/tree/dev"

UI_FIGURES = [
    ("fig01_table_list.png", "Список в виде таблицы"),
    ("fig02_card_list_mobile.png", "Список карточками на узком окне"),
    ("fig03_filters_applied.png", "Применённые фильтры отбора"),
    ("fig04_multi_selection.png", "Выделение нескольких записей"),
    ("fig05_empty_result.png", "Пустой результат поиска"),
    ("fig06_error_state.png", "Состояние ошибки загрузки"),
    ("fig07_url_query_params.png", "Адресная строка с параметрами"),
    ("fig08_equipment_detail.png", "Карточка оборудования"),
    ("fig09_loading_state.png", "Индикатор загрузки данных"),
]

CODE_FIGURES = [
    ("code_main.png", "Точка входа", "Router, provider и URL strategy"),
    ("code_equipment_model.png", "Модель Equipment", "Immutable сущность и copyWith"),
    ("code_client_model.png", "Модель Client", "Данные клиента проката"),
    ("code_page_result.png", "PageResult", "Пагинация и totalPages"),
    ("code_equipment_query.png", "EquipmentQuery", "Фильтры и сброс page"),
    ("code_client_query.png", "ClientQuery", "Условия отбора клиентов"),
    ("code_repo_equipment.png", "Репозиторий Equipment", "Поиск, фильтры, find"),
    ("code_repo_client.png", "Репозиторий Client", "In-memory выборка клиентов"),
    ("code_notifier_equipment.png", "Notifier оборудования", "Загрузка, query, selected"),
    ("code_notifier_client.png", "Notifier клиентов", "Состояние списка клиентов"),
    ("code_query_params.png", "Параметры URL", "Сериализация query в адрес"),
    ("code_entity_table.png", "EntityTable", "Обобщённая таблица сущностей"),
    ("code_screen_equipment_list.png", "Экран списка", "Таблица, фильтры, URL"),
    ("code_screen_equipment_detail.png", "Экран карточки", "Детальный просмотр Equipment"),
    ("code_screen_client_list.png", "Список клиентов", "Поиск и пагинация"),
    ("code_screen_client_detail.png", "Карточка клиента", "Просмотр полей Client"),
    ("code_pubspec.png", "Зависимости", "provider и go_router"),
    ("code_delete_many_fix.png", "Исправление deleteMany", "Проверка index >= 0"),
]


def _set_run_font(run, size: int = 14, bold: bool = False) -> None:
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run.bold = bold


def add_body(doc: Document, text: str, indent: float = 1.25) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.first_line_indent = Cm(indent)
    run = p.add_run(text)
    _set_run_font(run)


def add_heading(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.first_line_indent = Cm(0)
    run = p.add_run(text)
    _set_run_font(run, size=15, bold=True)


def add_page_break(doc: Document) -> None:
    p = doc.add_paragraph()
    run = p.add_run()
    br = OxmlElement("w:br")
    br.set(qn("w:type"), "page")
    run._r.append(br)


def add_center_caption(doc: Document, line: str, bold: bool = False) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    run = p.add_run(line)
    _set_run_font(run, size=12, bold=bold)


def add_figure(doc: Document, number: int, folder: Path, filename: str, caption: str) -> None:
    path = folder / filename
    if path.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Cm(0)
        run = p.add_run()
        run.add_picture(str(path), width=Cm(16))
    else:
        add_center_caption(doc, f"[Нет файла: {filename}]")
    add_center_caption(doc, f"Рисунок {number} — {caption}")


def add_code_figure(doc: Document, number: int, filename: str, title: str, desc: str) -> None:
    path = CODE_DIR / filename
    if path.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Cm(0)
        run = p.add_run()
        run.add_picture(str(path), width=Cm(15.5))
    add_center_caption(doc, f"Рисунок {number} — {title}", bold=True)
    add_center_caption(doc, desc)
    doc.add_paragraph()


def clone_title_page(doc: Document, template_path: Path) -> None:
    tpl = Document(str(template_path))
    body = doc.element.body
    for child in list(body):
        body.remove(child)
    for i, child in enumerate(tpl.element.body):
        if i > 25:
            break
        body.append(deepcopy(child))


def build_report() -> None:
    doc = Document()
    clone_title_page(doc, TITLE_TEMPLATE)
    add_page_break(doc)

    add_heading(doc, "Практическая работа 2. Списки, поиск, фильтрация и пагинация")
    add_body(
        doc,
        "Цель — разработать Flutter Web-каталог проката туристического и мототехнического "
        "оборудования с сущностями Equipment и Client, in-memory репозиториями, provider-состоянием, "
        "поиском, фильтрами, сортировкой, пагинацией и синхронизацией условий отбора с URL.",
    )

    add_heading(doc, "Схема слоёв приложения")
    add_body(
        doc,
        "Приложение разделено на четыре уровня: models (данные и Query), repositories (доступ к данным), "
        "state (ChangeNotifier через provider), screens и widgets (UI). Экраны не вызывают репозитории "
        "напрямую — только методы notifier. Такой порядок упрощает замену in-memory хранилища на HTTP API.",
    )
    add_figure(doc, 1, SCREENSHOTS, "fig_layers_schema.png", "Схема слоёв приложения")

    add_heading(doc, "Обобщённый виджет EntityTable")
    add_body(
        doc,
        "Таблица вынесена в EntityTable<T>, потому что Equipment и Client используют одинаковую механику: "
        "чекбоксы, сортировка по заголовку, прокрутка и колонка действий. Колонки задаются списком "
        "TableColumnSpec, поэтому при добавлении новой сущности не дублируется разметка DataTable — "
        "меняется только описание колонок.",
    )

    add_heading(doc, "Ошибка в методе deleteMany")
    add_body(
        doc,
        "В учебной версии deleteMany для Equipment использовалось условие index > 0 после indexWhere. "
        "Из-за этого элемент с индексом 0 в списке _items не удалялся при массовом удалении. "
        "Исправление — проверка index >= 0 (не найдено — index == -1).",
    )

    add_heading(doc, "Снимки экранов приложения")
    add_body(
        doc,
        "Ниже приведены обязательные состояния интерфейса: таблица, карточки на узком окне, фильтры, "
        "множественное выделение, пустой результат, ошибка, параметры в адресной строке, карточка и загрузка.",
    )

    fig = 2
    for fname, cap in UI_FIGURES:
        add_figure(doc, fig, SCREENSHOTS, fname, cap)
        fig += 1

    add_heading(doc, "Исходный код (ветка dev)")
    add_body(
        doc,
        f"Листинги приведены в виде изображений по состоянию репозитория {GITHUB_DEV}. "
        "Под каждым фрагментом — краткое название и назначение.",
    )

    for fname, title, desc in CODE_FIGURES:
        add_code_figure(doc, fig, fname, title, desc)
        fig += 1

    add_heading(doc, "Заключение")
    add_body(
        doc,
        "Реализованы каталоги оборудования и клиентов с полным набором операций списка, "
        "синхронизацией query-параметров URL, адаптивной вёрсткой и исправленной массовой "
        "удалением deleteMany. Проект проходит flutter analyze и готов к подключению backend.",
    )

    doc.save(OUTPUT)
    copy = ROOT / "Otchet_Praktika_2_Spiski_i_paginaciya.docx"
    doc.save(copy)
    print("Saved:", OUTPUT)


if __name__ == "__main__":
    build_report()
